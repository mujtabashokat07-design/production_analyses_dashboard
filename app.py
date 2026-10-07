from __future__ import annotations

from io import BytesIO
import json
import math
import re
import sqlite3
import unicodedata
import uuid
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st


WORKBOOK_PATH = Path(__file__).with_name("production_list_mill_1.xlsx")
DATABASE_PATH = Path(__file__).with_name("production_data.sqlite3")
HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1
SOURCE_COLUMNS = {
    "DATE": 1,
    "ITEM": 2,
    "GAUGE": 3,
    "QTY": 4,
    "WEIGHT": 5,
    "PATTA KG": 6,
    "BURADA SCRAP": 7,
    "PATRA SCRAP": 8,
    "PIPE SCRAP": 9,
    "TOTAL WEIGHT": 10,
}
NUMERIC_COLUMNS = (
    "QTY",
    "WEIGHT",
    "PATTA KG",
    "BURADA SCRAP",
    "PATRA SCRAP",
    "PIPE SCRAP",
    "TOTAL WEIGHT",
)
SCRAP_COLUMNS = ("BURADA SCRAP", "PATRA SCRAP", "PIPE SCRAP")
CHART_COLORS = ["#40C9FF", "#8B7CFF", "#41D6A3", "#FFB86B", "#FF6B87"]
DARK_PLOT = {
    "paper_bgcolor": "rgba(0,0,0,0)",
    "plot_bgcolor": "rgba(0,0,0,0)",
    "font": {"color": "#D7E1EE", "family": "Inter, Segoe UI, sans-serif"},
    "xaxis": {
        "gridcolor": "rgba(150,170,195,0.12)",
        "linecolor": "rgba(150,170,195,0.18)",
        "zerolinecolor": "rgba(150,170,195,0.18)",
    },
    "yaxis": {
        "gridcolor": "rgba(150,170,195,0.12)",
        "linecolor": "rgba(150,170,195,0.18)",
        "zerolinecolor": "rgba(150,170,195,0.18)",
    },
    "margin": {"l": 12, "r": 16, "t": 34, "b": 12},
    "legend": {"orientation": "h", "y": 1.12, "x": 0},
}


def normalize_item_label(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = re.sub(r"\s*[xX×�]\s*", " × ", text)
    text = re.sub(r'\s*"\s*', '"', text)
    text = re.sub(r"\s*/\s*", "/", text)
    text = re.sub(r"(?<=\d)\s*rd\b", " rd", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_item_identity(value: Any) -> str:
    return normalize_item_label(value).casefold()


def normalize_gauge_label(value: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip().casefold()


def parse_source_date(value: Any, excel_row: int) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        text = value.strip()
        for date_format in ("%d/%m/%Y", "%Y-%m-%d", "%m-%d-%y"):
            try:
                return datetime.strptime(text, date_format)
            except ValueError:
                continue
    raise ValueError(
        f"Could not interpret the date in worksheet row {excel_row}: {value!r}"
    )


def source_date_display(value: Any, number_format: str) -> str:
    if isinstance(value, (datetime, date)):
        date_value = value if isinstance(value, datetime) else datetime.combine(
            value, datetime.min.time()
        )
        if number_format.lower() == "mm-dd-yy":
            return date_value.strftime("%m-%d-%y")
        if number_format.lower() == "dd/mm/yyyy":
            return date_value.strftime("%d/%m/%Y")
        return date_value.strftime("%Y-%m-%d")
    return str(value)


def numeric_value(value: Any, column: str, excel_row: int) -> tuple[float | None, str | None]:
    if value is None or value == "":
        return None, None
    try:
        return float(value), None
    except (TypeError, ValueError):
        return None, f"Row {excel_row}: {column} contains {value!r}"


@st.cache_data(show_spinner=False)
def load_production_data(
    workbook_path: str, modified_time_ns: int
) -> tuple[pd.DataFrame, list[str], int]:
    del modified_time_ns
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        worksheet = workbook.active
        records: list[dict[str, Any]] = []
        data_issues: list[str] = []
        excluded_rows = 0

        for excel_row in range(FIRST_DATA_ROW, worksheet.max_row + 1):
            raw_values = {
                name: worksheet.cell(excel_row, column).value
                for name, column in SOURCE_COLUMNS.items()
            }
            raw_date = raw_values["DATE"]
            if raw_date is None or raw_date == "":
                if any(value is not None for value in raw_values.values()):
                    excluded_rows += 1
                continue

            item = raw_values["ITEM"]
            if item is None or not str(item).strip():
                data_issues.append(
                    f"Row {excel_row}: date exists but item is blank; row omitted from production records."
                )
                continue

            date_cell = worksheet.cell(excel_row, SOURCE_COLUMNS["DATE"])
            parsed_date = parse_source_date(raw_date, excel_row)
            record: dict[str, Any] = {
                "Excel row": excel_row,
                "_record_id": f"excel:{excel_row}",
                "_date": parsed_date,
                "DATE": source_date_display(raw_date, date_cell.number_format),
                "ITEM": item,
                "GAUGE": raw_values["GAUGE"],
                "_item_group": normalize_item_label(item),
                "_gauge_group": (
                    normalize_gauge_label(raw_values["GAUGE"])
                    if raw_values["GAUGE"] is not None
                    else None
                ),
            }
            for column in NUMERIC_COLUMNS:
                value, issue = numeric_value(raw_values[column], column, excel_row)
                record[column] = raw_values[column]
                record[f"_num_{column}"] = value
                if issue:
                    data_issues.append(issue)
            records.append(record)

        return pd.DataFrame.from_records(records), data_issues, excluded_rows
    finally:
        workbook.close()


def initialize_production_store(database_path: Path) -> None:
    with closing(sqlite3.connect(database_path)) as connection:
        with connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS production_records (
                    record_id TEXT PRIMARY KEY,
                    production_date TEXT NOT NULL,
                    item TEXT NOT NULL,
                    gauge TEXT,
                    qty REAL,
                    weight REAL,
                    patta_kg REAL,
                    burada_scrap REAL,
                    patra_scrap REAL,
                    pipe_scrap REAL,
                    total_weight REAL
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS deleted_source_records (
                    record_id TEXT PRIMARY KEY
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS source_record_overrides (
                    record_id TEXT PRIMARY KEY,
                    payload TEXT NOT NULL
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS sales_records (
                    sale_id TEXT PRIMARY KEY,
                    sale_date TEXT NOT NULL,
                    item TEXT NOT NULL,
                    item_key TEXT NOT NULL,
                    gauge TEXT,
                    gauge_key TEXT NOT NULL,
                    quantity REAL NOT NULL CHECK (quantity > 0)
                )
                """
            )
            connection.execute(
                """
                CREATE INDEX IF NOT EXISTS sales_records_stock_key
                ON sales_records (item_key, gauge_key)
                """
            )


def load_source_record_overrides(database_path: Path) -> dict[str, dict[str, Any]]:
    initialize_production_store(database_path)
    with closing(sqlite3.connect(database_path)) as connection:
        rows = connection.execute(
            "SELECT record_id, payload FROM source_record_overrides"
        )
        return {
            str(record_id): json.loads(payload)
            for record_id, payload in rows
        }


def save_source_record_overrides(
    database_path: Path, overrides: dict[str, dict[str, Any]]
) -> None:
    initialize_production_store(database_path)
    with closing(sqlite3.connect(database_path)) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for record_id, payload in overrides.items():
            existing_row = connection.execute(
                "SELECT payload FROM source_record_overrides WHERE record_id = ?",
                (record_id,),
            ).fetchone()
            merged_payload = (
                json.loads(existing_row[0]) if existing_row is not None else {}
            )
            merged_payload.update(payload)
            connection.execute(
                """
                INSERT INTO source_record_overrides (record_id, payload)
                VALUES (?, ?)
                ON CONFLICT(record_id) DO UPDATE SET payload = excluded.payload
                """,
                (record_id, json.dumps(merged_payload, allow_nan=False)),
            )
        connection.commit()


def apply_source_record_overrides(
    source_data: pd.DataFrame,
    overrides: dict[str, dict[str, Any]],
) -> pd.DataFrame:
    if source_data.empty or not overrides:
        return source_data

    updated = source_data.copy()
    row_indexes = {
        str(record_id): index
        for index, record_id in updated["_record_id"].items()
    }
    for record_id, payload in overrides.items():
        if record_id not in row_indexes:
            continue
        index = row_indexes[record_id]
        for column, value in payload.items():
            if column == "DATE":
                parsed_date = parse_source_date(value, int(updated.at[index, "Excel row"]))
                updated.at[index, "_date"] = parsed_date
                updated.at[index, "DATE"] = parsed_date.strftime("%Y-%m-%d")
            elif column == "ITEM":
                item = str(value).strip()
                updated.at[index, "ITEM"] = item
                updated.at[index, "_item_group"] = normalize_item_label(item)
            elif column == "GAUGE":
                gauge = str(value).strip() or None
                updated.at[index, "GAUGE"] = gauge
                updated.at[index, "_gauge_group"] = (
                    normalize_gauge_label(gauge) if gauge else None
                )
            elif column in NUMERIC_COLUMNS:
                numeric = None if value is None else float(value)
                updated.at[index, column] = numeric
                updated.at[index, f"_num_{column}"] = numeric
            else:
                raise ValueError(f"Unsupported source record field in saved edit: {column}")
    return updated


def reset_dashboard_data(database_path: Path) -> None:
    initialize_production_store(database_path)
    with closing(sqlite3.connect(database_path)) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for table in (
            "production_records",
            "sales_records",
            "deleted_source_records",
            "source_record_overrides",
        ):
            connection.execute(f"DELETE FROM {table}")
        connection.commit()


def load_saved_production_data(
    database_path: Path,
) -> tuple[list[dict[str, Any]], set[str]]:
    initialize_production_store(database_path)
    with closing(sqlite3.connect(database_path)) as connection:
        connection.row_factory = sqlite3.Row
        records = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM production_records ORDER BY production_date DESC, record_id"
            )
        ]
        deleted_source_ids = {
            row["record_id"]
            for row in connection.execute("SELECT record_id FROM deleted_source_records")
        }
    return records, deleted_source_ids


def save_production_record(database_path: Path, record: dict[str, Any]) -> None:
    with closing(sqlite3.connect(database_path)) as connection:
        with connection:
            connection.execute(
                """
                INSERT INTO production_records (
                    record_id, production_date, item, gauge, qty, weight, patta_kg,
                    burada_scrap, patra_scrap, pipe_scrap, total_weight
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    record["record_id"],
                    record["production_date"],
                    record["item"],
                    record["gauge"],
                    record["QTY"],
                    record["WEIGHT"],
                    record["PATTA KG"],
                    record["BURADA SCRAP"],
                    record["PATRA SCRAP"],
                    record["PIPE SCRAP"],
                    record["TOTAL WEIGHT"],
                ),
            )


def load_sales_records(database_path: Path) -> list[dict[str, Any]]:
    initialize_production_store(database_path)
    with closing(sqlite3.connect(database_path)) as connection:
        connection.row_factory = sqlite3.Row
        return [
            dict(row)
            for row in connection.execute(
                """
                SELECT sale_id, sale_date, item, item_key, gauge, gauge_key, quantity
                FROM sales_records
                ORDER BY sale_date DESC, sale_id DESC
                """
            )
        ]


def build_sales_inventory(
    production_data: pd.DataFrame, sales_records: list[dict[str, Any]]
) -> pd.DataFrame:
    inventory_columns = [
        "_item_group",
        "_gauge_key",
        "ITEM",
        "GAUGE",
        "PRODUCED QTY",
        "PRODUCED WEIGHT",
        "SOLD QTY",
        "AVAILABLE QTY",
        "STOCK SHORTFALL",
    ]
    if production_data.empty:
        return pd.DataFrame(columns=inventory_columns)

    stock = production_data.copy()
    stock["_item_group"] = stock["ITEM"].map(normalize_item_identity)
    stock["_gauge_key"] = stock["GAUGE"].map(
        lambda value: normalize_gauge_label(value) if pd.notna(value) else ""
    )
    stock["_gauge_label"] = stock["GAUGE"].map(
        lambda value: str(value).strip() if pd.notna(value) else ""
    )
    stock["_item_label"] = stock["ITEM"].map(normalize_item_label)
    stock_by_variant = (
        stock.groupby(["_item_group", "_gauge_key"], as_index=False, dropna=False)
        .agg(
            ITEM=("_item_label", "first"),
            GAUGE=(
                "_gauge_label",
                lambda values: next((value for value in values if value), ""),
            ),
            **{
                "PRODUCED QTY": ("_num_QTY", lambda values: values.sum(min_count=1)),
                "PRODUCED WEIGHT": (
                    "_num_WEIGHT",
                    lambda values: values.sum(min_count=1),
                ),
            },
        )
    )
    stock_by_variant["PRODUCED QTY"] = stock_by_variant["PRODUCED QTY"].fillna(0.0)
    stock_by_variant["PRODUCED WEIGHT"] = stock_by_variant[
        "PRODUCED WEIGHT"
    ].fillna(0.0)

    sales = pd.DataFrame.from_records(sales_records)
    if sales.empty:
        stock_by_variant["SOLD QTY"] = 0.0
    else:
        sales["_item_group"] = sales["item"].map(normalize_item_identity)
        sales["_gauge_key"] = sales["gauge"].map(
            lambda value: normalize_gauge_label(value) if pd.notna(value) else ""
        )
        sold_by_variant = sales.groupby(
            ["_item_group", "_gauge_key"], as_index=False
        )["quantity"].sum()
        sold_by_variant = sold_by_variant.rename(
            columns={"quantity": "SOLD QTY"}
        )
        stock_by_variant = stock_by_variant.merge(
            sold_by_variant,
            how="left",
            on=["_item_group", "_gauge_key"],
            validate="one_to_one",
        )
        stock_by_variant["SOLD QTY"] = stock_by_variant["SOLD QTY"].fillna(0.0)

    stock_by_variant["STOCK SHORTFALL"] = (
        stock_by_variant["SOLD QTY"] - stock_by_variant["PRODUCED QTY"]
    ).clip(lower=0.0)
    stock_by_variant["AVAILABLE QTY"] = (
        stock_by_variant["PRODUCED QTY"] - stock_by_variant["SOLD QTY"]
    ).clip(lower=0.0)
    return stock_by_variant[inventory_columns].sort_values(
        ["ITEM", "GAUGE"], key=lambda values: values.astype("string").str.casefold()
    ).reset_index(drop=True)


def record_sale(
    database_path: Path,
    sale: dict[str, Any],
    produced_quantity: float,
) -> None:
    initialize_production_store(database_path)
    quantity = float(sale["quantity"])
    if not math.isfinite(quantity) or quantity <= 0:
        raise ValueError("Sale quantity must be a positive number.")
    if not math.isclose(quantity, round(quantity), abs_tol=1e-9):
        raise ValueError("Sale quantity must be a whole number of pieces.")
    if not math.isfinite(produced_quantity) or produced_quantity < 0:
        raise ValueError("Available production quantity is invalid.")
    item = sale["item"].strip()
    item_key = sale["item_key"].strip()
    if not item or not item_key or normalize_item_identity(item) != item_key:
        raise ValueError("Select a valid item before recording a sale.")
    if not isinstance(sale["sale_date"], str) or not sale["sale_date"]:
        raise ValueError("Select a valid sale date.")
    try:
        date.fromisoformat(sale["sale_date"])
    except ValueError as error:
        raise ValueError("Sale date must be a valid calendar date.") from error
    gauge = sale["gauge"]
    gauge_key = sale["gauge_key"]
    expected_gauge_key = normalize_gauge_label(gauge) if gauge else ""
    if gauge_key != expected_gauge_key:
        raise ValueError("Selected gauge does not match the item variant.")

    with closing(sqlite3.connect(database_path)) as connection:
        connection.execute("BEGIN IMMEDIATE")
        existing_sales = connection.execute(
            "SELECT item, gauge, quantity FROM sales_records"
        )
        previously_sold = sum(
            float(existing_quantity)
            for existing_item, existing_gauge, existing_quantity in existing_sales
            if normalize_item_identity(existing_item) == item_key
            and (
                normalize_gauge_label(existing_gauge)
                if existing_gauge
                else ""
            )
            == gauge_key
        )
        available_quantity = max(0.0, produced_quantity - float(previously_sold))
        if quantity > available_quantity + 1e-9:
            raise ValueError(
                f"Only {available_quantity:,.2f} pcs are currently available for this item and gauge."
            )
        connection.execute(
            """
            INSERT INTO sales_records (
                sale_id, sale_date, item, item_key, gauge, gauge_key, quantity
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sale["sale_id"],
                sale["sale_date"],
                item,
                item_key,
                gauge,
                gauge_key,
                quantity,
            ),
        )
        connection.commit()


def number_label(value: float, decimals: int = 0) -> str:
    return f"{value:,.{decimals}f}"


def style_chart(fig: go.Figure, height: int = 330) -> go.Figure:
    fig.update_layout(**DARK_PLOT, height=height, hoverlabel={"bgcolor": "#182333"})
    fig.update_xaxes(automargin=True)
    fig.update_yaxes(automargin=True)
    return fig


st.set_page_config(
    page_title="Engineering Zone | Production Analysis Dashboard",
    page_icon="🏭",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    :root {
        --surface: #101925;
        --surface-raised: #151f2d;
        --line: rgba(157, 178, 204, .15);
        --muted: #91a2b8;
        --text: #e8eef6;
        --accent: #40c9ff;
    }
    .stApp {
        background:
            radial-gradient(ellipse at 4% -10%, rgba(35, 111, 157, .24), transparent 34%),
            radial-gradient(ellipse at 96% 0%, rgba(93, 76, 163, .18), transparent 30%),
            linear-gradient(180deg, #0b111a 0%, #0d1520 48%, #0b111a 100%);
        color: var(--text);
    }
    [data-testid="stHeader"] { background: rgba(11, 17, 26, .82); }
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #111b28 0%, #0e1722 100%);
        border-right: 1px solid var(--line);
    }
    [data-testid="stSidebar"] > div:first-child { padding-top: 1.4rem; }
    .block-container { padding-top: 1.5rem; padding-bottom: 3rem; max-width: 1580px; }
    .eyebrow {
        color: #69d4ff; font-size: .72rem; font-weight: 750;
        letter-spacing: .16em; text-transform: uppercase; margin-bottom: .45rem;
    }
    .hero-title { font-size: clamp(1.8rem, 3vw, 2.65rem); line-height: 1.08; font-weight: 780; margin: 0; letter-spacing: -.045em; }
    .hero-subtitle { color: #91a2b8; margin-top: .55rem; font-size: .98rem; max-width: 760px; }
    [data-testid="stSegmentedControl"] {
        background: rgba(17, 27, 40, .78); border: 1px solid var(--line);
        border-radius: 14px; padding: .3rem; margin: .8rem 0 1.4rem;
    }
    [data-testid="stSegmentedControl"] button {
        min-height: 2.7rem; border-radius: 10px;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: linear-gradient(145deg, rgba(19, 30, 44, .9), rgba(13, 21, 32, .92));
        border-color: var(--line); border-radius: 17px;
    }
    [data-testid="stBaseButton-primary"] {
        border-radius: 10px; min-height: 2.8rem; font-weight: 700;
        box-shadow: 0 8px 22px rgba(64, 201, 255, .14);
    }
    .metric-card {
        background: linear-gradient(145deg, rgba(23, 35, 50, .98), rgba(17, 26, 38, .98));
        border: 1px solid var(--line); border-radius: 16px;
        padding: 1rem .95rem; min-height: 128px;
        box-shadow: 0 10px 32px rgba(0, 0, 0, .18);
    }
    .metric-label {
        color: #9aaabd; font-size: .76rem; font-weight: 650;
        letter-spacing: .06em; text-transform: uppercase;
    }
    .metric-value {
        color: #f3f7fc; font-size: clamp(1.05rem, 1.4vw, 1.35rem);
        font-weight: 740; letter-spacing: -.035em; white-space: nowrap; margin-top: .43rem;
    }
    .metric-note { color: #8394aa; font-size: .78rem; margin-top: .25rem; }
    .section-heading { font-size: 1.13rem; font-weight: 710; margin: .2rem 0 .1rem; letter-spacing: -.015em; }
    .section-caption { color: #8fa0b5; font-size: .84rem; margin-bottom: .8rem; }
    .insight-card {
        background: rgba(21, 31, 45, .78); border: 1px solid var(--line);
        border-radius: 13px; padding: .95rem 1rem; height: 100%;
    }
    .insight-label { color: #92a4b9; font-size: .75rem; text-transform: uppercase; letter-spacing: .08em; }
    .insight-value { color: #edf4fb; font-size: 1rem; font-weight: 650; margin-top: .35rem; }
    div[data-testid="stDataFrame"] { border: 1px solid var(--line); border-radius: 12px; overflow: hidden; }
    [data-baseweb="tag"] {
        background-color: #183b52 !important; color: #e8f7ff !important;
    }
    [data-baseweb="select"] > div {
        background-color: #111b28; border-color: var(--line);
    }
    hr { border-color: var(--line); }
    @media (max-width: 700px) {
        .block-container { padding-top: 1rem; }
        .metric-value { white-space: normal; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

try:
    if WORKBOOK_PATH.is_file():
        workbook_mtime = WORKBOOK_PATH.stat().st_mtime_ns
        data, data_issues, excluded_rows = load_production_data(
            str(WORKBOOK_PATH), workbook_mtime
        )
    else:
        data = pd.DataFrame(
            columns=[
                "Excel row",
                "_record_id",
                "_date",
                "DATE",
                "ITEM",
                "GAUGE",
                "_item_group",
                "_gauge_group",
                *NUMERIC_COLUMNS,
                *(f"_num_{column}" for column in NUMERIC_COLUMNS),
            ]
        )
        data_issues = []
        excluded_rows = 0
    saved_records, deleted_source_ids = load_saved_production_data(DATABASE_PATH)
    sales_records = load_sales_records(DATABASE_PATH)
    source_overrides = load_source_record_overrides(DATABASE_PATH)
except (
    OSError,
    ValueError,
    sqlite3.Error,
    openpyxl.utils.exceptions.InvalidFileException,
) as error:
    st.error(f"Unable to load production data or the local data store: {error}")
    st.stop()

try:
    data = apply_source_record_overrides(data, source_overrides)
except (TypeError, ValueError, KeyError) as error:
    st.error(f"Unable to apply saved source-row edits: {error}")
    st.stop()

if deleted_source_ids and "_record_id" in data.columns:
    data = data.loc[~data["_record_id"].isin(deleted_source_ids)].copy()

if saved_records:
    added_records: list[dict[str, Any]] = []
    for saved_record in saved_records:
        record_date = datetime.strptime(
            saved_record["production_date"], "%Y-%m-%d"
        )
        record: dict[str, Any] = {
            "Excel row": None,
            "_record_id": saved_record["record_id"],
            "_date": record_date,
            "DATE": record_date.strftime("%Y-%m-%d"),
            "ITEM": saved_record["item"],
            "GAUGE": saved_record["gauge"],
            "_item_group": normalize_item_label(saved_record["item"]),
            "_gauge_group": (
                normalize_gauge_label(saved_record["gauge"])
                if saved_record["gauge"]
                else None
            ),
        }
        for column, database_column in zip(
            NUMERIC_COLUMNS,
            (
                "qty",
                "weight",
                "patta_kg",
                "burada_scrap",
                "patra_scrap",
                "pipe_scrap",
                "total_weight",
            ),
        ):
            value = saved_record[database_column]
            record[column] = value
            record[f"_num_{column}"] = value
        added_records.append(record)
    data = pd.concat(
        [data, pd.DataFrame.from_records(added_records)], ignore_index=True
    )

if data.empty:
    data = pd.DataFrame(
        columns=[
            "Excel row",
            "_record_id",
            "_date",
            "DATE",
            "ITEM",
            "GAUGE",
            "_item_group",
            "_gauge_group",
            *NUMERIC_COLUMNS,
            *(f"_num_{column}" for column in NUMERIC_COLUMNS),
        ]
    )
else:
    for column in NUMERIC_COLUMNS:
        numeric_column = f"_num_{column}"
        data[numeric_column] = pd.to_numeric(data[numeric_column], errors="coerce")

st.markdown('<div class="eyebrow">ENGINEERING ZONE&nbsp;&nbsp; / &nbsp;&nbsp;OPERATIONS</div>', unsafe_allow_html=True)
st.markdown('<h1 class="hero-title">Production intelligence</h1>', unsafe_allow_html=True)
st.markdown(
    '<div class="hero-subtitle">Production performance and live data entry in one clear workspace.</div>',
    unsafe_allow_html=True,
)
if st.session_state.get("workspace_view") not in {"Overview", "Add", "Sales & Stock"}:
    st.session_state["workspace_view"] = "Overview"
active_view = st.segmented_control(
    "Workspace",
    ("Overview", "Add", "Sales & Stock"),
    key="workspace_view",
    label_visibility="collapsed",
    width="stretch",
)

if active_view == "Add":
    st.subheader("Add production records", icon=":material/add_circle:")
    st.caption(
        "Choose an existing item and gauge or enter a new value. Each save stays in production history; "
        "matching item and gauge quantities are combined automatically in current stock."
    )
    with st.container(border=True):
        item_suggestions = sorted(
            {str(value).strip() for value in data["ITEM"].dropna() if str(value).strip()},
            key=str.casefold,
        )

        def reset_numeric_suggestions() -> None:
            for label in NUMERIC_COLUMNS:
                suggestion_key = f"new_record_suggestion_{label}"
                if st.session_state.get(suggestion_key) != "Enter manually":
                    st.session_state[suggestion_key] = "Enter manually"
                    st.session_state[f"new_record_custom_{label}"] = None

        def reset_item_dependent_suggestions() -> None:
            gauge_key = "new_record_gauge_choice"
            if st.session_state.get(gauge_key) not in {
                "No gauge",
                "Enter a new gauge",
            }:
                st.session_state[gauge_key] = "No gauge"
            reset_numeric_suggestions()

        item_entry_col, gauge_entry_col = st.columns(2)
        with item_entry_col:
            item_choice = st.selectbox(
                "Item suggestion",
                ["Enter a new item", *item_suggestions],
                key="new_record_item_choice",
                on_change=reset_item_dependent_suggestions,
                help="Search existing items or select the first option to type a new item.",
            )
            entry_item = (
                st.text_input("New item name", max_chars=120, key="new_record_item")
                if item_choice == "Enter a new item"
                else item_choice
            )

        if item_choice == "Enter a new item":
            selected_item_data = data.iloc[0:0]
        else:
            selected_item_data = data.loc[
                data["ITEM"].map(normalize_item_identity)
                == normalize_item_identity(item_choice)
            ]
        gauge_suggestions = sorted(
            {
                str(value).strip()
                for value in selected_item_data["GAUGE"].dropna()
                if str(value).strip()
            },
            key=str.casefold,
        )
        with gauge_entry_col:
            gauge_choice = st.selectbox(
                "Gauge suggestion",
                ["No gauge", "Enter a new gauge", *gauge_suggestions],
                key="new_record_gauge_choice",
                on_change=reset_numeric_suggestions,
                help="Choose a gauge previously used for this item, enter a new one, or leave it blank.",
            )
            entry_gauge = (
                st.text_input("New gauge", max_chars=80, key="new_record_gauge")
                if gauge_choice == "Enter a new gauge"
                else ("" if gauge_choice == "No gauge" else gauge_choice)
            )

        if gauge_choice == "Enter a new gauge":
            selected_variant_data = data.iloc[0:0]
        elif gauge_choice == "No gauge":
            blank_gauge = selected_item_data["GAUGE"].isna() | (
                selected_item_data["GAUGE"].astype("string").str.strip() == ""
            )
            selected_variant_data = selected_item_data.loc[blank_gauge]
        else:
            selected_gauge_key = normalize_gauge_label(gauge_choice)
            selected_variant_data = selected_item_data.loc[
                selected_item_data["GAUGE"].map(
                    lambda value: (
                        normalize_gauge_label(value) if pd.notna(value) else ""
                    )
                )
                == selected_gauge_key
            ]

        numeric_suggestions: dict[str, list[float]] = {}
        for label in NUMERIC_COLUMNS:
            values = pd.to_numeric(
                selected_variant_data[f"_num_{label}"], errors="coerce"
            ).dropna()
            common_values = values.value_counts().head(12).index.astype(float).tolist()
            recent_values = (
                pd.to_numeric(
                    selected_variant_data[f"_num_{label}"].tail(12), errors="coerce"
                )
                .dropna()
                .astype(float)
                .tolist()
            )
            numeric_suggestions[label] = sorted(set(common_values + recent_values))

        st.markdown("#### Recent production values")
        st.caption(
            "Type any quantity, weight, or scrap value directly. Choose a recent value above an input to fill it quickly; leave the input blank if it was not recorded."
        )

        def apply_numeric_suggestion(label: str) -> None:
            selected_value = st.session_state[f"new_record_suggestion_{label}"]
            st.session_state[f"new_record_custom_{label}"] = (
                None
                if selected_value == "Enter manually"
                else float(selected_value)
            )

        def mark_numeric_value_manual(label: str) -> None:
            st.session_state[f"new_record_suggestion_{label}"] = "Enter manually"

        numeric_inputs: dict[str, float | None] = {}
        for labels in (NUMERIC_COLUMNS[:4], NUMERIC_COLUMNS[4:]):
            numeric_layout = st.columns(len(labels))
            for column, label in zip(numeric_layout, labels):
                with column:
                    st.selectbox(
                        f"Recent {label.lower()} values",
                        ["Enter manually", *numeric_suggestions[label]],
                        key=f"new_record_suggestion_{label}",
                        format_func=lambda option: (
                            f"{option:,.2f}" if isinstance(option, float) else option
                        ),
                        on_change=apply_numeric_suggestion,
                        args=(label,),
                        help="Choose a previous value to fill the editable input below.",
                    )
                    value = st.number_input(
                        label,
                        min_value=0.0,
                        value=None,
                        step=1.0,
                        format="%.2f",
                        key=f"new_record_custom_{label}",
                        on_change=mark_numeric_value_manual,
                        args=(label,),
                        help="Enter a new value here, or edit a suggested value.",
                    )
                    numeric_inputs[label] = None if value is None else float(value)

        entry_date = st.date_input(
            "Production date", value=date.today(), format="DD/MM/YYYY"
        )
        with st.form("add_production_record"):
            submitted = st.form_submit_button(
                "Save production record",
                type="primary",
                icon=":material/save:",
                width="stretch",
            )

    if submitted:
        if not entry_item.strip():
            st.error("Enter an item name before saving.")
        else:
            new_record = {
                "record_id": f"manual:{uuid.uuid4().hex}",
                "production_date": entry_date.isoformat(),
                "item": entry_item.strip(),
                "gauge": entry_gauge.strip() or None,
                **numeric_inputs,
            }
            try:
                save_production_record(DATABASE_PATH, new_record)
            except sqlite3.Error as error:
                st.error(f"Unable to save the production record: {error}")
            else:
                st.toast("Production record saved. Opening the updated overview.", icon=":material/check_circle:")
                st.rerun()
    st.stop()

if active_view == "Sales & Stock":
    st.subheader("Sales & stock", icon=":material/point_of_sale:")
    st.caption(
        "Select a produced item and gauge, record the pieces sold, and the remaining stock updates immediately. "
        "Production history stays unchanged; weight is shown as a reference only."
    )
    inventory = build_sales_inventory(data, sales_records)
    available_stock = inventory.loc[inventory["AVAILABLE QTY"] >= 1].copy()
    total_available = float(inventory["AVAILABLE QTY"].sum())
    total_sold = sum(float(sale["quantity"]) for sale in sales_records)
    stock_items = int(
        inventory.loc[inventory["AVAILABLE QTY"] >= 1, "_item_group"].nunique()
    )
    stock_metrics = st.columns(4)
    stock_metrics[0].metric("Items in stock", f"{stock_items:,}", border=True)
    stock_metrics[1].metric(
        "Available variants", f"{len(available_stock):,}", border=True
    )
    stock_metrics[2].metric("Pieces sold", f"{total_sold:,.0f}", border=True)
    stock_metrics[3].metric("Pieces remaining", f"{total_available:,.0f}", border=True)

    if inventory["STOCK SHORTFALL"].gt(1e-9).any():
        st.warning(
            "Some saved sales exceed the production quantity currently in the register. "
            "Review the production source before recording more sales for those variants."
        )

    st.markdown("#### Available stock")
    if available_stock.empty:
        st.info("There are no unsold pieces available to sell.")
    else:
        stock_table = available_stock[
            [
                "ITEM",
                "GAUGE",
                "PRODUCED QTY",
                "SOLD QTY",
                "AVAILABLE QTY",
                "PRODUCED WEIGHT",
            ]
        ].rename(
            columns={
                "ITEM": "Item",
                "GAUGE": "Gauge",
                "PRODUCED QTY": "Produced (pcs)",
                "SOLD QTY": "Sold (pcs)",
                "AVAILABLE QTY": "Available (pcs)",
                "PRODUCED WEIGHT": "Produced weight (kg, reference)",
            }
        )
        stock_table["Gauge"] = stock_table["Gauge"].replace("", "No gauge")
        st.dataframe(stock_table, width="stretch", hide_index=True)

    if available_stock.empty:
        st.info("Add production records before recording a sale.")
    else:
        with st.container(border=True):
            st.markdown("#### Record a sale")
            item_labels = (
                available_stock.drop_duplicates("_item_group")
                .set_index("_item_group")["ITEM"]
                .to_dict()
            )
            selected_item_key = st.selectbox(
                "Select item",
                options=list(item_labels),
                format_func=lambda item_key: item_labels[item_key],
                key="sale_item",
            )
            item_variants = available_stock.loc[
                available_stock["_item_group"] == selected_item_key
            ]
            gauge_labels = item_variants.set_index("_gauge_key")["GAUGE"].to_dict()
            selected_gauge_key = st.selectbox(
                "Select gauge",
                options=list(gauge_labels),
                format_func=lambda gauge_key: gauge_labels[gauge_key] or "No gauge",
                key=f"sale_gauge_{selected_item_key}",
            )
            selected_stock = item_variants.loc[
                item_variants["_gauge_key"] == selected_gauge_key
            ].iloc[0]
            available_quantity = float(selected_stock["AVAILABLE QTY"])
            reference_weight = float(selected_stock["PRODUCED WEIGHT"])
            stock_detail_columns = st.columns(2)
            stock_detail_columns[0].metric(
                "Available quantity", f"{available_quantity:,.0f} pcs"
            )
            stock_detail_columns[1].metric(
                "Produced weight (reference only)", f"{reference_weight:,.2f} kg"
            )
            maximum_sale_quantity = math.floor(available_quantity + 1e-9)
            if maximum_sale_quantity < 1:
                st.info(
                    "Less than one whole piece remains. The sales register records pieces as whole units."
                )
            else:
                with st.form("record_sale", clear_on_submit=True):
                    sale_date = st.date_input(
                        "Sale date",
                        value=date.today(),
                        format="DD/MM/YYYY",
                        key=f"sale_date_{selected_item_key}_{selected_gauge_key}",
                    )
                    sale_quantity = st.number_input(
                        "Quantity sold (pcs)",
                        min_value=1,
                        max_value=maximum_sale_quantity,
                        value=1,
                        step=1,
                        key=f"sale_quantity_{selected_item_key}_{selected_gauge_key}",
                        help=f"Up to {maximum_sale_quantity:,} pieces are available for this item and gauge.",
                    )
                    sale_submitted = st.form_submit_button(
                        "Record sale and update stock",
                        type="primary",
                        icon=":material/point_of_sale:",
                        width="stretch",
                    )
                if sale_submitted:
                    sale = {
                        "sale_id": f"sale:{uuid.uuid4().hex}",
                        "sale_date": sale_date.isoformat(),
                        "item": str(selected_stock["ITEM"]),
                        "item_key": str(selected_stock["_item_group"]),
                        "gauge": (
                            None
                            if not str(selected_stock["GAUGE"]).strip()
                            else str(selected_stock["GAUGE"])
                        ),
                        "gauge_key": str(selected_stock["_gauge_key"]),
                        "quantity": float(sale_quantity),
                    }
                    try:
                        record_sale(
                            DATABASE_PATH,
                            sale,
                            float(selected_stock["PRODUCED QTY"]),
                        )
                    except (sqlite3.Error, ValueError) as error:
                        st.error(f"Unable to record the sale: {error}")
                    else:
                        st.toast(
                            f"Sale recorded: {sale_quantity:,} pcs. Available stock updated.",
                            icon=":material/check_circle:",
                        )
                        st.rerun()

    st.markdown("#### Recent sales")
    if not sales_records:
        st.info("No sales have been recorded yet.")
    else:
        recent_sales = pd.DataFrame.from_records(sales_records[:100])[
            ["sale_date", "item", "gauge", "quantity"]
        ].rename(
            columns={
                "sale_date": "Sale date",
                "item": "Item",
                "gauge": "Gauge",
                "quantity": "Quantity sold (pcs)",
            }
        )
        recent_sales["Item"] = recent_sales["Item"].map(normalize_item_label)
        recent_sales["Gauge"] = recent_sales["Gauge"].fillna("No gauge")
        st.caption(
            f"Showing the latest {len(recent_sales):,} of {len(sales_records):,} recorded sales."
        )
        st.dataframe(recent_sales, width="stretch", hide_index=True)
    st.stop()

st.sidebar.markdown('<div class="eyebrow">MILL 1 • OPERATIONS</div>', unsafe_allow_html=True)
st.sidebar.title("Dashboard filters")
st.sidebar.caption("Filters apply to all metrics and visualizations.")

if data.empty:
    st.sidebar.info("No production records are currently available.")
    start_date = end_date = date.today()
else:
    min_date = data["_date"].min().date()
    max_date = data["_date"].max().date()
    selected_dates = st.sidebar.date_input(
        "Production date range",
        value=(min_date, max_date),
        min_value=min_date,
        max_value=max_date,
        format="DD/MM/YYYY",
    )
    if isinstance(selected_dates, tuple) and len(selected_dates) == 2:
        start_date, end_date = selected_dates
    elif isinstance(selected_dates, tuple) and len(selected_dates) == 1:
        start_date = end_date = selected_dates[0]
    else:
        start_date = end_date = selected_dates

item_options = sorted(data["_item_group"].dropna().unique().tolist(), key=str.casefold)
gauge_options = sorted(data["_gauge_group"].dropna().unique().tolist(), key=str.casefold)
st.sidebar.caption("Leave item and gauge filters empty to include all values.")
selected_items = st.sidebar.multiselect("Item", item_options, placeholder="All items")
selected_gauges = st.sidebar.multiselect("Gauge", gauge_options, placeholder="All gauges")

with st.sidebar.expander("Reset dashboard data"):
    st.caption(
        "Removes manually added production and sales, restores deleted Excel rows, "
        "and clears saved Excel-row edits. The original workbook stays unchanged."
    )
    with st.form("reset_dashboard_data", clear_on_submit=True):
        confirm_reset = st.checkbox("I understand this cannot be undone.")
        reset_submitted = st.form_submit_button(
            "Reset to original Excel data",
            type="secondary",
            width="stretch",
        )
    if reset_submitted:
        if not confirm_reset:
            st.sidebar.error("Confirm the reset before continuing.")
        else:
            try:
                reset_dashboard_data(DATABASE_PATH)
            except sqlite3.Error as error:
                st.sidebar.error(f"Unable to reset dashboard data: {error}")
            else:
                st.cache_data.clear()
                st.rerun()

period = st.sidebar.radio("Trend interval", ("Monthly", "Daily"), horizontal=True)
if data.empty:
    filtered = data.copy()
else:
    item_filter = data["_item_group"].isin(selected_items) if selected_items else True
    gauge_filter = data["_gauge_group"].isin(selected_gauges) if selected_gauges else True
    filtered = data.loc[
        data["_date"].dt.date.between(start_date, end_date)
        & item_filter
        & gauge_filter
    ].copy()

overview_inventory = build_sales_inventory(data, sales_records)
if selected_items:
    selected_item_keys = {
        normalize_item_identity(item) for item in selected_items
    }
    overview_inventory = overview_inventory.loc[
        overview_inventory["_item_group"].isin(selected_item_keys)
    ].copy()
if selected_gauges:
    overview_inventory = overview_inventory.loc[
        overview_inventory["_gauge_key"].isin(selected_gauges)
    ].copy()
current_sold_quantity = float(overview_inventory["SOLD QTY"].sum())
current_available_quantity = float(overview_inventory["AVAILABLE QTY"].sum())

st.markdown('<div class="eyebrow">PRODUCTION INTELLIGENCE</div>', unsafe_allow_html=True)
st.markdown('<h2 class="section-heading">Production overview</h2>', unsafe_allow_html=True)
st.markdown(
    '<div class="section-caption">A live view of output, product mix and recorded scrap for the selected filters.</div>',
    unsafe_allow_html=True,
)

if filtered.empty:
    st.info(
        "No production records match the selected filters. Adjust the filters or add a new record above."
    )
    st.stop()

total_weight = filtered["_num_TOTAL WEIGHT"].sum(min_count=1)
total_quantity = filtered["_num_QTY"].sum(min_count=1)
total_scrap = filtered[[f"_num_{column}" for column in SCRAP_COLUMNS]].sum().sum()
patta_measurement = filtered["_num_PATTA KG"].sum()
total_weight_value = float(total_weight) if pd.notna(total_weight) else 0.0
total_quantity_value = float(total_quantity) if pd.notna(total_quantity) else 0.0
record_count = len(filtered)
scrap_rate = (total_scrap / total_weight_value * 100) if total_weight_value else 0.0
invalid_qty = int((filtered["QTY"].notna() & filtered["_num_QTY"].isna()).sum())

metric_specs = [
    ("TOTAL WEIGHT", f"{number_label(total_weight_value)} kg", "Sum of the source TOTAL WEIGHT field"),
    ("PATTA KG • WEIGHING", f"{number_label(float(patta_measurement))} kg", "Separate weighing measurement — not scrap"),
    (
        "PRODUCTION QTY",
        f"{number_label(total_quantity_value)} pcs",
        f"Produced within the selected date range{f' • {invalid_qty} invalid' if invalid_qty else ''}",
    ),
    ("PRODUCTION RECORDS", number_label(record_count), "Rows matching the current filters"),
    ("RECORDED SCRAP", f"{number_label(float(total_scrap))} kg", "Burada + Patra + Pipe scrap"),
    ("SCRAP / TOTAL WEIGHT", f"{scrap_rate:.2f}%", "Descriptive ratio of recorded scrap to total weight"),
    (
        "SOLD QTY",
        f"{number_label(current_sold_quantity)} pcs",
        "Recorded sales for the selected item and gauge filters",
    ),
    (
        "AVAILABLE STOCK",
        f"{number_label(current_available_quantity)} pcs",
        "Current produced stock minus recorded sales",
    ),
]
for row_start in range(0, len(metric_specs), 3):
    if row_start:
        st.write("")
    metric_cols = st.columns(3)
    for column, (label, value, note) in zip(
        metric_cols, metric_specs[row_start : row_start + 3]
    ):
        with column:
            st.markdown(
                f'<div class="metric-card"><div class="metric-label">{label}</div>'
                f'<div class="metric-value">{value}</div>'
                f'<div class="metric-note">{note}</div></div>',
                unsafe_allow_html=True,
            )

st.write("")
st.markdown(
    '<div class="section-heading">Current stock by item and gauge</div>',
    unsafe_allow_html=True,
)
st.markdown(
    '<div class="section-caption">Production quantities remain in the production analysis. '
    'Available stock below deducts saved sales and combines matching item and gauge entries; '
    'the production date filter does not change current stock.</div>',
    unsafe_allow_html=True,
)
overview_stock_table = overview_inventory[
    [
        "ITEM",
        "GAUGE",
        "PRODUCED QTY",
        "SOLD QTY",
        "AVAILABLE QTY",
        "PRODUCED WEIGHT",
    ]
].rename(
    columns={
        "ITEM": "Item",
        "GAUGE": "Gauge",
        "PRODUCED QTY": "Produced (pcs)",
        "SOLD QTY": "Sold (pcs)",
        "AVAILABLE QTY": "Available (pcs)",
        "PRODUCED WEIGHT": "Produced weight (kg, reference)",
    }
)
overview_stock_table["Gauge"] = overview_stock_table["Gauge"].replace("", "No gauge")
st.dataframe(
    overview_stock_table,
    width="stretch",
    hide_index=True,
    height=360,
)

st.write("")
monthly = (
    filtered.assign(_period=filtered["_date"].dt.to_period("M").dt.to_timestamp())
    .groupby("_period", as_index=False)
    .agg(
        **{
            "TOTAL WEIGHT": ("_num_TOTAL WEIGHT", "sum"),
            "QTY": ("_num_QTY", "sum"),
        }
    )
)
if period == "Daily":
    trend = (
        filtered.assign(_period=filtered["_date"].dt.normalize())
        .groupby("_period", as_index=False)
        .agg(
            **{
                "TOTAL WEIGHT": ("_num_TOTAL WEIGHT", "sum"),
                "QTY": ("_num_QTY", "sum"),
            }
        )
    )
else:
    trend = monthly

top_items = (
    filtered.groupby("_item_group", as_index=False)
    .agg(
        **{
            "TOTAL WEIGHT": ("_num_TOTAL WEIGHT", "sum"),
            "QTY": ("_num_QTY", "sum"),
        }
    )
    .rename(columns={"_item_group": "ITEM"})
    .nlargest(10, "TOTAL WEIGHT")
    .sort_values("TOTAL WEIGHT")
)
trend_col, item_col = st.columns([1.55, 1], gap="large")
with trend_col:
    st.markdown('<div class="section-heading">Output over time</div>', unsafe_allow_html=True)
    st.markdown(
        f'<div class="section-caption">{period} output: total weight (kg) and quantity (pcs).</div>',
        unsafe_allow_html=True,
    )
    trend_fig = go.Figure()
    trend_fig.add_trace(
        go.Scatter(
            x=trend["_period"],
            y=trend["TOTAL WEIGHT"],
            name="Total weight (kg)",
            mode="lines+markers",
            line={"color": CHART_COLORS[0], "width": 3},
            marker={"size": 7},
            hovertemplate="%{x|%d %b %Y}<br>Total weight: %{y:,.0f} kg<extra></extra>",
        )
    )
    trend_fig.add_trace(
        go.Bar(
            x=trend["_period"],
            y=trend["QTY"],
            name="Quantity (pcs)",
            marker={"color": "rgba(139,124,255,.52)"},
            yaxis="y2",
            hovertemplate="%{x|%d %b %Y}<br>Quantity: %{y:,.0f} pcs<extra></extra>",
        )
    )
    trend_fig.update_layout(
        **{key: value for key, value in DARK_PLOT.items() if key != "yaxis"},
        height=350,
        hovermode="x unified",
        yaxis={
            "title": "Total weight (kg)",
            "gridcolor": "rgba(150,170,195,0.12)",
            "linecolor": "rgba(150,170,195,0.18)",
            "zerolinecolor": "rgba(150,170,195,0.18)",
        },
        yaxis2={
            "title": "Quantity (pcs)",
            "overlaying": "y",
            "side": "right",
            "showgrid": False,
        },
        barmode="overlay",
    )
    st.plotly_chart(trend_fig, width="stretch")

with item_col:
    st.markdown('<div class="section-heading">Top items by total weight (kg)</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-caption">Highest-output items in the filtered dataset.</div>',
        unsafe_allow_html=True,
    )
    item_fig = px.bar(
        top_items,
        x="TOTAL WEIGHT",
        y="ITEM",
        orientation="h",
        text="TOTAL WEIGHT",
        color_discrete_sequence=[CHART_COLORS[0]],
    )
    item_fig.update_traces(
        texttemplate="%{x:,.0f} kg",
        textposition="outside",
        cliponaxis=False,
        hovertemplate="%{y}<br>Total weight: %{x:,.0f} kg<extra></extra>",
    )
    item_fig.update_layout(
        showlegend=False,
        xaxis_title="Total weight (kg)",
        yaxis_title="",
        xaxis={"range": [0, float(top_items["TOTAL WEIGHT"].max()) * 1.28]},
    )
    st.plotly_chart(style_chart(item_fig, 350), width="stretch")

st.write("")
scrap_col, gauge_col, insight_col = st.columns([1, 1.25, 1], gap="large")
with scrap_col:
    st.markdown('<div class="section-heading">Recorded scrap mix</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-caption">Scrap amounts are in kg. Patta KG is a separate weighing measurement.</div>',
        unsafe_allow_html=True,
    )
    scrap_data = pd.DataFrame(
        {
            "Category": list(SCRAP_COLUMNS),
            "Recorded amount": [
                float(filtered[f"_num_{column}"].sum()) for column in SCRAP_COLUMNS
            ],
        }
    )
    scrap_data = scrap_data.loc[scrap_data["Recorded amount"] > 0]
    scrap_fig = px.pie(
        scrap_data,
        names="Category",
        values="Recorded amount",
        hole=0.68,
        color_discrete_sequence=CHART_COLORS,
    )
    scrap_fig.update_traces(
        textinfo="percent",
        textfont_size=12,
        hovertemplate="%{label}<br>Amount: %{value:,.0f} kg<br>Share: %{percent}<extra></extra>",
    )
    scrap_fig.update_layout(showlegend=True, legend={"orientation": "v", "x": 1.02, "y": 0.5})
    st.plotly_chart(style_chart(scrap_fig, 330), width="stretch")

with gauge_col:
    st.markdown('<div class="section-heading">Output by gauge</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-caption">Top 10 gauge groups ranked by total weight (kg).</div>',
        unsafe_allow_html=True,
    )
    gauge_data = (
        filtered.groupby("_gauge_group", as_index=False)
        .agg(**{"TOTAL WEIGHT": ("_num_TOTAL WEIGHT", "sum")})
        .rename(columns={"_gauge_group": "GAUGE"})
        .nlargest(10, "TOTAL WEIGHT")
        .sort_values("TOTAL WEIGHT")
    )
    gauge_fig = px.bar(
        gauge_data,
        x="TOTAL WEIGHT",
        y="GAUGE",
        orientation="h",
        color_discrete_sequence=[CHART_COLORS[1]],
        text="TOTAL WEIGHT",
    )
    gauge_fig.update_traces(
        texttemplate="%{x:,.0f} kg",
        textposition="outside",
        cliponaxis=False,
        hovertemplate="%{y}<br>Total weight: %{x:,.0f} kg<extra></extra>",
    )
    gauge_fig.update_layout(
        showlegend=False,
        xaxis_title="Total weight (kg)",
        yaxis_title="",
        xaxis={"range": [0, float(gauge_data["TOTAL WEIGHT"].max()) * 1.28]},
    )
    st.plotly_chart(style_chart(gauge_fig, 330), width="stretch")

with insight_col:
    st.markdown('<div class="section-heading">At a glance</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-caption">Highlights for the current filter selection.</div>',
        unsafe_allow_html=True,
    )
    best_item = top_items.iloc[-1]["ITEM"] if not top_items.empty else "—"
    best_month = monthly.loc[monthly["TOTAL WEIGHT"].idxmax()] if not monthly.empty else None
    leading_scrap = scrap_data.loc[scrap_data["Recorded amount"].idxmax()] if not scrap_data.empty else None
    insights = [
        ("Leading item", str(best_item)),
        (
            "Peak output month",
            best_month["_period"].strftime("%b %Y") if best_month is not None else "—",
        ),
        (
            "Largest scrap category",
            str(leading_scrap["Category"]) if leading_scrap is not None else "No recorded scrap",
        ),
        (
            "Distinct items / gauges",
            f"{filtered['_item_group'].nunique()} / {filtered['_gauge_group'].nunique()}",
        ),
    ]
    for label, value in insights:
        st.markdown(
            f'<div class="insight-card" style="margin-bottom:.58rem">'
            f'<div class="insight-label">{label}</div>'
            f'<div class="insight-value">{value}</div></div>',
            unsafe_allow_html=True,
        )

st.write("")
st.markdown('<div class="section-heading">Production records</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="section-caption">Excel records retain their source row number; saved entries '
    'remain available after app restarts.</div>',
    unsafe_allow_html=True,
)
display_columns = [
    "_record_id",
    "Excel row",
    "DATE",
    "ITEM",
    "GAUGE",
    "QTY",
    "WEIGHT",
    "PATTA KG",
    "BURADA SCRAP",
    "PATRA SCRAP",
    "PIPE SCRAP",
    "TOTAL WEIGHT",
            ]
display_data = (
    filtered[["_date", *display_columns]]
    .sort_values(["_date", "_record_id"], ascending=[False, False])
    .drop(columns=["_date"])
    .rename(columns={"_record_id": "Record ID"})
)
display_data["Record ID"] = display_data["Record ID"].map(
    lambda record_id: (
        f"Manual {record_id.removeprefix('manual:')[:8]}"
        if record_id.startswith("manual:")
        else record_id
    )
)
for column in NUMERIC_COLUMNS:
    has_non_numeric_source_value = any(
        not pd.isna(value) and not isinstance(value, (int, float))
        for value in display_data[column]
    )
    if has_non_numeric_source_value:
        display_data[column] = display_data[column].map(
            lambda value: "" if pd.isna(value) else str(value)
        )
display_data = display_data.rename(
    columns={
        "QTY": "QTY (pcs)",
        "WEIGHT": "WEIGHT (kg)",
        "PATTA KG": "PATTA KG (weighing)",
        "BURADA SCRAP": "BURADA SCRAP (kg)",
        "PATRA SCRAP": "PATRA SCRAP (kg)",
        "PIPE SCRAP": "PIPE SCRAP (kg)",
        "TOTAL WEIGHT": "TOTAL WEIGHT (kg)",
    }
)

with st.expander("Download selected production rows", expanded=False):
    st.caption(
        "Choose rows from the current Overview filters, then choose exactly which columns to include."
    )
    export_row_data = filtered.loc[display_data.index]
    export_row_labels: list[str] = []
    export_record_ids: dict[str, str] = {}
    for _, row in export_row_data.iterrows():
        record_id = str(row["_record_id"])
        excel_row = row["Excel row"]
        row_label = (
            f"Excel row {int(excel_row)}"
            if pd.notna(excel_row)
            else f"Manual record {record_id.removeprefix('manual:')[:8]}"
        )
        gauge_label = str(row["GAUGE"]).strip() if pd.notna(row["GAUGE"]) else ""
        row_label += f" | {row['ITEM']} | {gauge_label or 'No gauge'} | {row['DATE']}"
        if row_label in export_record_ids:
            row_label += f" | {record_id}"
        export_row_labels.append(row_label)
        export_record_ids[row_label] = record_id

    selected_export_rows = st.multiselect(
        "Rows to export",
        options=export_row_labels,
        default=export_row_labels,
        key="overview_export_rows",
        placeholder="Select production rows",
    )
    available_export_columns = display_data.columns.tolist()
    selected_export_columns = st.multiselect(
        "Columns to export",
        options=available_export_columns,
        default=available_export_columns,
        key="overview_export_columns",
        placeholder="Select data columns",
    )
    selected_record_ids = {
        export_record_ids[label] for label in selected_export_rows
    }
    selected_export_data = display_data.loc[
        export_row_data["_record_id"].astype(str).isin(selected_record_ids),
        selected_export_columns,
    ]
    export_buffer = BytesIO()
    with pd.ExcelWriter(export_buffer, engine="openpyxl") as writer:
        selected_export_data.to_excel(
            writer,
            index=False,
            sheet_name="Production Data",
        )
    st.download_button(
        "Download selected rows and columns",
        data=export_buffer.getvalue(),
        file_name="production_data.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        disabled=not selected_export_rows or not selected_export_columns,
        type="primary",
    )

with st.expander("Edit Excel source rows", expanded=False):
    st.caption(
        "Edits are saved in the dashboard database only; the original Excel workbook is not changed. "
        "Use the date, item, and gauge filters to find the rows you want to edit."
    )
    editable_source = filtered.loc[
        filtered["_record_id"].astype(str).str.startswith("excel:")
    ].copy()
    if editable_source.empty:
        st.info("No Excel source rows match the current Overview filters.")
    else:
        editor_columns = [
            "_record_id",
            "Excel row",
            "DATE",
            "ITEM",
            "GAUGE",
            *NUMERIC_COLUMNS,
        ]
        editor_data = pd.DataFrame(
            {
                "_record_id": editable_source["_record_id"].astype(str),
                "Excel row": editable_source["Excel row"].astype(int),
                "DATE": editable_source["_date"].dt.date,
                "ITEM": editable_source["ITEM"].astype(str),
                "GAUGE": editable_source["GAUGE"].map(
                    lambda value: "" if pd.isna(value) else str(value)
                ),
                **{
                    column: editable_source[f"_num_{column}"]
                    for column in NUMERIC_COLUMNS
                },
            },
            columns=editor_columns,
        )
        editor_key = (
            "overview_source_editor_"
            + str(hash(tuple(editor_data["_record_id"].tolist())))
        )
        with st.form(f"save_source_edits_{editor_key}"):
            edited_source = st.data_editor(
                editor_data,
                key=editor_key,
                width="stretch",
                hide_index=True,
                num_rows="fixed",
                column_config={
                    "_record_id": st.column_config.TextColumn(
                        "Record ID", disabled=True
                    ),
                    "Excel row": st.column_config.NumberColumn(
                        "Excel row", disabled=True, format="%d"
                    ),
                    "DATE": st.column_config.DateColumn(
                        "Date", format="DD/MM/YYYY", required=True
                    ),
                    "ITEM": st.column_config.TextColumn(
                        "Item", required=True
                    ),
                    "GAUGE": st.column_config.TextColumn("Gauge"),
                    **{
                        column: st.column_config.NumberColumn(
                            column, format="%.2f", step=0.01
                        )
                        for column in NUMERIC_COLUMNS
                    },
                },
            )
            save_source_edits = st.form_submit_button(
                "Save Excel row edits",
                type="primary",
                icon=":material/save:",
            )

        if save_source_edits:
            changed_overrides: dict[str, dict[str, Any]] = {}
            validation_errors: list[str] = []
            source_by_id = editable_source.set_index("_record_id", drop=False)
            for _, edited_row in edited_source.iterrows():
                record_id = str(edited_row["_record_id"])
                original = source_by_id.loc[record_id]
                row_number = int(original["Excel row"])
                changes: dict[str, Any] = {}
                try:
                    edited_date = parse_source_date(
                        edited_row["DATE"], row_number
                    ).strftime("%Y-%m-%d")
                except (TypeError, ValueError):
                    validation_errors.append(
                        f"Excel row {row_number}: enter a valid production date."
                    )
                    continue
                original_date = original["_date"].strftime("%Y-%m-%d")
                if edited_date != original_date:
                    changes["DATE"] = edited_date

                edited_item = (
                    "" if pd.isna(edited_row["ITEM"]) else str(edited_row["ITEM"]).strip()
                )
                if not edited_item:
                    validation_errors.append(
                        f"Excel row {row_number}: item cannot be blank."
                    )
                    continue
                if edited_item != str(original["ITEM"]).strip():
                    changes["ITEM"] = edited_item

                edited_gauge = (
                    "" if pd.isna(edited_row["GAUGE"]) else str(edited_row["GAUGE"]).strip()
                )
                original_gauge = (
                    "" if pd.isna(original["GAUGE"]) else str(original["GAUGE"]).strip()
                )
                if edited_gauge != original_gauge:
                    changes["GAUGE"] = edited_gauge

                for column in NUMERIC_COLUMNS:
                    edited_value = edited_row[column]
                    numeric_value = (
                        None if pd.isna(edited_value) else float(edited_value)
                    )
                    original_value = original[f"_num_{column}"]
                    original_numeric = (
                        None if pd.isna(original_value) else float(original_value)
                    )
                    if numeric_value != original_numeric:
                        changes[column] = numeric_value
                if changes:
                    changed_overrides[record_id] = changes

            if validation_errors:
                st.error(" ".join(validation_errors))
            elif not changed_overrides:
                st.info("No Excel source-row changes to save.")
            else:
                try:
                    save_source_record_overrides(
                        DATABASE_PATH,
                        changed_overrides,
                    )
                except (sqlite3.Error, TypeError, ValueError) as error:
                    st.error(f"Unable to save Excel row edits: {error}")
                else:
                    st.toast(
                        f"Saved edits for {len(changed_overrides):,} Excel row(s).",
                        icon=":material/check_circle:",
                    )
                    st.rerun()

st.dataframe(
    display_data,
    width="stretch",
    hide_index=True,
    height=440,
)

quality_notes: list[str] = []
if data_issues:
    quality_notes.append(
        f"{len(data_issues)} source value(s) need attention; non-numeric values are excluded from numeric sums."
    )
if excluded_rows:
    quality_notes.append(
        f"{excluded_rows} non-record row(s) without a date were excluded (such as worksheet totals)."
    )
if quality_notes:
    with st.expander("Data quality and calculation notes", expanded=bool(data_issues)):
        for note in quality_notes:
            st.warning(note)
        for issue in data_issues:
            st.caption(issue)
        st.caption(
            "Blank scrap cells are excluded from the reported scrap sums. PATTA KG is a "
            "separate weighing measurement and is not classified as scrap. Only BURADA SCRAP, "
            "PATRA SCRAP, and PIPE SCRAP are counted as scrap. No source cells, "
            "date values, or workbook formatting are edited."
        )
        st.caption(
            "The source TOTAL WEIGHT field is displayed as recorded; the dashboard does not "
            "recalculate it from the separate weighing and scrap columns."
        )
else:
    st.caption(
        "PATTA KG is a separate weighing measurement, not scrap. Scrap totals include only "
        "BURADA SCRAP, PATRA SCRAP, and PIPE SCRAP. The source workbook is read-only; "
        "new records are saved in the local data store."
    )
